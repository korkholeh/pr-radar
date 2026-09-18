"""T16: `exports/xlsx.py::write_xlsx()`, read back with openpyxl."""

from __future__ import annotations

import datetime
import io

import openpyxl
import pytest

from apps.dashboards.exports.columns import ExportColumn
from apps.dashboards.exports.xlsx import write_xlsx

COLUMNS = [
    ExportColumn(key="name", title="Name", type="url", link_key="url", width=30),
    ExportColumn(key="count", title="Count", type="int"),
    ExportColumn(key="ratio", title="Ratio", type="percent"),
    ExportColumn(key="dur", title="Duration", type="duration"),
    ExportColumn(key="opened", title="Opened", type="datetime"),
    ExportColumn(key="formula_looking", title="=cmd()", type="text"),
]

ROW = {
    "name": "PR #1",
    "url": "https://github.com/example/repo/pull/1",
    "count": 3,
    "ratio": 0.457,
    "dur": 7200.0,
    "opened": datetime.datetime(2026, 8, 1, 12, 0, tzinfo=datetime.UTC),
    "formula_looking": "text",
}
EMPTY_ROW = {
    "name": None,
    "url": None,
    "count": None,
    "ratio": None,
    "dur": None,
    "opened": None,
    "formula_looking": None,
}


def _workbook(rows):
    return openpyxl.load_workbook(io.BytesIO(write_xlsx(COLUMNS, rows, "Test Sheet")))


def test_workbook_formatting() -> None:
    workbook = _workbook([ROW])
    sheet = workbook.active
    header_cell = sheet.cell(row=1, column=1)
    assert header_cell.font.bold is True
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F2"


def test_numeric_percent_and_duration_cells() -> None:
    workbook = _workbook([ROW])
    sheet = workbook.active
    count_cell = sheet.cell(row=2, column=2)
    assert count_cell.value == 3
    assert isinstance(count_cell.value, int | float)

    ratio_cell = sheet.cell(row=2, column=3)
    assert ratio_cell.value == pytest.approx(0.457)
    assert ratio_cell.number_format == "0.0%"

    duration_cell = sheet.cell(row=2, column=4)
    assert duration_cell.value == pytest.approx(2.0)
    assert duration_cell.number_format == "0.0"

    duration_header = sheet.cell(row=1, column=4)
    assert duration_header.value == "Duration, h"


def test_pr_url_is_a_hyperlink() -> None:
    workbook = _workbook([ROW])
    sheet = workbook.active
    name_cell = sheet.cell(row=2, column=1)
    assert name_cell.value == "PR #1"
    assert name_cell.hyperlink is not None
    assert name_cell.hyperlink.target == "https://github.com/example/repo/pull/1"


def test_none_is_an_empty_cell() -> None:
    workbook = _workbook([EMPTY_ROW])
    sheet = workbook.active
    for column_index in range(1, len(COLUMNS) + 1):
        assert sheet.cell(row=2, column=column_index).value is None


def test_delta_cells_use_their_own_numeric_format_not_percent() -> None:
    """Round-1 BLOCKER regression: a delta column typed `int`/`duration` (per the
    `exports/columns.py::_metric_columns` fix) must keep that number format, not the `0.0%`
    format a hardcoded `type="percent"` delta column used to apply to every metric's delta."""
    columns = [
        ExportColumn(key="prs_merged__delta", title="PRs merged (Δ)", type="int"),
        ExportColumn(key="lead_time_p50__delta", title="Lead time (Δ)", type="duration"),
    ]
    workbook = openpyxl.load_workbook(
        io.BytesIO(
            write_xlsx(columns, [{"prs_merged__delta": 5.0, "lead_time_p50__delta": -3600.0}], "Sheet")
        )
    )
    sheet = workbook.active
    count_delta_cell = sheet.cell(row=2, column=1)
    assert count_delta_cell.value == 5
    assert count_delta_cell.number_format != "0.0%"

    duration_delta_cell = sheet.cell(row=2, column=2)
    assert duration_delta_cell.value == pytest.approx(-1.0)
    assert duration_delta_cell.number_format == "0.0"


def test_delta_cells_are_coloured_green_or_red_by_metric_direction() -> None:
    """Round 2 review MINOR fix (spec §10.6): a delta column carries `MetricDef.direction`, and
    `write_xlsx()` colours the cell green when that delta is an improvement, red when it's a
    regression — independent of whether the delta is positive or negative in absolute terms."""
    columns = [
        ExportColumn(
            key="prs_merged__delta", title="PRs merged (Δ)", type="int", direction="higher_is_better"
        ),
        ExportColumn(
            key="lead_time_p50__delta", title="Lead time (Δ)", type="duration", direction="lower_is_better"
        ),
        ExportColumn(key="pr_size_p50__delta", title="PR size (Δ)", type="float", direction="neutral"),
    ]
    row = {
        "prs_merged__delta": -3.0,  # higher_is_better, negative -> regression -> red
        "lead_time_p50__delta": -3600.0,  # lower_is_better, negative -> improvement -> green
        "pr_size_p50__delta": 2.0,  # neutral -> never coloured
    }
    workbook = openpyxl.load_workbook(io.BytesIO(write_xlsx(columns, [row], "Sheet")))
    sheet = workbook.active

    prs_merged_cell = sheet.cell(row=2, column=1)
    lead_time_cell = sheet.cell(row=2, column=2)
    pr_size_cell = sheet.cell(row=2, column=3)

    assert prs_merged_cell.fill.fgColor.rgb == "FFFBE9E7"
    assert lead_time_cell.fill.fgColor.rgb == "FFE6F4EA"
    assert pr_size_cell.fill.fgColor.rgb in (None, "00000000")


def test_a_zero_delta_is_never_coloured() -> None:
    columns = [
        ExportColumn(
            key="prs_merged__delta", title="PRs merged (Δ)", type="int", direction="higher_is_better"
        )
    ]
    workbook = openpyxl.load_workbook(io.BytesIO(write_xlsx(columns, [{"prs_merged__delta": 0.0}], "Sheet")))
    sheet = workbook.active
    cell = sheet.cell(row=2, column=1)
    assert cell.fill.fgColor.rgb in (None, "00000000")


def test_formula_looking_title_is_stored_as_a_string_not_a_formula() -> None:
    workbook = _workbook([ROW])
    sheet = workbook.active
    header_cell = sheet.cell(row=1, column=6)
    assert header_cell.value == "=cmd()"
    assert header_cell.data_type == "s"
