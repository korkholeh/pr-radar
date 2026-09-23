"""T17/T18: `exports/reports.py::build_report()` (plan §5, spec §10.6, acceptance criterion #7).

Design choices this test verifies (logged as unconfirmed in DECISIONS after session 2, now proven
or fixed here): sheet names/order per level, the level's own table sheets only (`recent_prs`
excluded since the report's `PRs` sheet already covers every PR), native Excel charts on Trends,
a Summary value equal to a direct `compute()` call, and the Parameters sheet's fields.

Policy is out of scope here: it is not a `Scope`/`DashboardParams` level (T19's job to wire), so
the "five levels" of the plan's acceptance table narrows to the four `ScopeType` values."""

from __future__ import annotations

import datetime
import io

import openpyxl
import pytest

from apps.accounts.selectors import ScopeFilter
from apps.activity.factories import PullRequestFactory
from apps.catalog.factories import IdentityFactory, PersonFactory, ProjectFactory, RepositoryFactory
from apps.dashboards.exports.reports import build_report, report_filename
from apps.dashboards.params import DashboardParams
from apps.dashboards.pr_filters import PRFilters
from apps.metrics.models import ScopeType
from apps.metrics.rollups import rebuild
from apps.metrics.services import bump_data_version, compute
from apps.metrics.types import Scope
from apps.policy.factories import PolicyViolationFactory

DATE_FROM = datetime.date(2026, 8, 1)
DATE_TO = datetime.date(2026, 8, 31)

_EXPECTED_TABLE_SHEETS = {
    ScopeType.GLOBAL: ("Projects", "People"),
    ScopeType.PROJECT: ("Repositories", "People"),
    ScopeType.REPO: ("People",),
    ScopeType.PERSON: (),
}


def _params(*, mode: str = "period", day: datetime.date = DATE_TO) -> DashboardParams:
    return DashboardParams(
        mode=mode,
        preset="custom",
        date_from=DATE_FROM,
        date_to=DATE_TO,
        day=day,
        granularity="week",
        granularity_is_auto=False,
        cohort="all",
        project_ids=(),
        repository_ids=(),
        q="",
        sort="",
        page=1,
        table="",
        pr_filters=PRFilters(),
    )


@pytest.fixture
def seeded_org(db):
    """One project, one repository, two people each with a merged PR in the period, and one open
    violation — enough for every sheet (Summary, Trends, tables, PRs, Violations) to have data."""
    repository = RepositoryFactory()
    project = ProjectFactory()
    project.repositories.add(repository)
    people = []
    for _ in range(2):
        person = PersonFactory()
        identity = IdentityFactory(person=person)
        pr = PullRequestFactory(
            repository=repository,
            author=identity,
            state="merged",
            created_at=datetime.datetime(2026, 8, 5, tzinfo=datetime.UTC),
            merged_at=datetime.datetime(2026, 8, 10, tzinfo=datetime.UTC),
            last_activity_at=datetime.datetime(2026, 8, 10, tzinfo=datetime.UTC),
        )
        people.append((person, pr))
    PolicyViolationFactory(pull_request=people[0][1])
    rebuild(DATE_FROM, DATE_TO)
    bump_data_version()
    return {"project": project, "repository": repository, "people": people}


def _workbook(scope: Scope, params: DashboardParams, *, user) -> openpyxl.Workbook:
    content = build_report(scope, params, user=user, language="en")
    return openpyxl.load_workbook(io.BytesIO(content))


@pytest.mark.django_db
@pytest.mark.parametrize("mode", ["period", "day"])
@pytest.mark.parametrize(
    "scope_type", [ScopeType.GLOBAL, ScopeType.PROJECT, ScopeType.REPO, ScopeType.PERSON]
)
def test_report_contains_every_mandated_sheet(scope_type, mode, seeded_org, django_user_model):
    project = seeded_org["project"]
    repository = seeded_org["repository"]
    person = seeded_org["people"][0][0]
    scope_id = {
        ScopeType.GLOBAL: None,
        ScopeType.PROJECT: project.id,
        ScopeType.REPO: repository.id,
        ScopeType.PERSON: person.id,
    }[scope_type]
    scope = Scope(scope_type=scope_type, scope_id=scope_id, access=ScopeFilter(unrestricted=True))
    user = django_user_model.objects.create_user(username="lead", password="pw")

    workbook = _workbook(scope, _params(mode=mode), user=user)

    expected = (
        ("Summary", "Trends")
        + _EXPECTED_TABLE_SHEETS[scope_type]
        + ("PRs", "Violations", "Metrics", "Parameters")
    )
    assert tuple(workbook.sheetnames) == expected


@pytest.mark.django_db
def test_trends_sheet_has_native_excel_charts(seeded_org, django_user_model):
    scope = Scope(scope_type=ScopeType.GLOBAL, scope_id=None, access=ScopeFilter(unrestricted=True))
    user = django_user_model.objects.create_user(username="lead", password="pw")

    workbook = _workbook(scope, _params(), user=user)

    trends_sheet = workbook["Trends"]
    assert trends_sheet.max_row > 1  # header + at least one data row
    assert len(trends_sheet._charts) == 3


@pytest.mark.django_db
def test_trends_sheet_follows_the_charts_it_is_built_from(seeded_org, django_user_model):
    """The sheet pivots the Throughput, AI adoption and Latency chart payloads by index, so a series
    dropped from a chart (the Latency p50 lines, `disclosure_rate`) must leave the sheet too — or the
    report fails on a missing dataset."""
    scope = Scope(scope_type=ScopeType.GLOBAL, scope_id=None, access=ScopeFilter(unrestricted=True))
    user = django_user_model.objects.create_user(username="lead", password="pw")

    trends_sheet = _workbook(scope, _params(), user=user)["Trends"]

    header = [cell.value for cell in trends_sheet[1]]
    assert len(header) == 6
    assert not any("p50" in str(title) or "isclosure" in str(title) for title in header)


@pytest.mark.django_db
def test_summary_value_matches_direct_compute_call(seeded_org, django_user_model):
    scope = Scope(scope_type=ScopeType.GLOBAL, scope_id=None, access=ScopeFilter(unrestricted=True))
    params = _params()
    user = django_user_model.objects.create_user(username="lead", password="pw")

    from apps.dashboards.exports.reports import _summary_metric_keys

    metric_keys = _summary_metric_keys(scope, params)
    assert metric_keys, "fixture must seed at least one summary metric"
    first_key = metric_keys[0]
    expected = compute([first_key], scope, params.date_from, params.date_to, cohort="all", granularity="week")
    expected_value = expected[first_key].value

    workbook = _workbook(scope, params, user=user)
    summary_sheet = workbook["Summary"]
    value_cell = summary_sheet.cell(row=2, column=2).value  # first metric row, "Value" column

    assert value_cell == expected_value


@pytest.mark.django_db
def test_parameters_sheet_has_the_documented_fields(seeded_org, django_user_model):
    scope = Scope(scope_type=ScopeType.GLOBAL, scope_id=None, access=ScopeFilter(unrestricted=True))
    user = django_user_model.objects.create_user(username="lead", password="pw")

    workbook = _workbook(scope, _params(), user=user)
    parameters_sheet = workbook["Parameters"]
    names = {parameters_sheet.cell(row=r, column=1).value for r in range(2, parameters_sheet.max_row + 1)}

    for expected_name in (
        "User",
        "Generated at",
        "Last successful sync",
        "Tool version",
        "Report timezone",
        "Report language",
    ):
        assert expected_name in names

    username_row = next(
        r
        for r in range(2, parameters_sheet.max_row + 1)
        if parameters_sheet.cell(row=r, column=1).value == "User"
    )
    assert parameters_sheet.cell(row=username_row, column=2).value == user.get_username()


@pytest.mark.django_db
def test_violations_sheet_is_narrowed_to_the_report_scope(seeded_org, django_user_model):
    """Round 1 review MAJOR: `_violation_rows()` used to read `violations_in_scope(scope.access)`,
    which applies only the *access* filter, not the page's own `Scope` — a project- or
    person-scoped report listed every visible project's violations. It must now go through
    `metrics.selectors.scoped_violations(scope)`, same as the PRs sheet and every other table."""

    def _violation_in_period(pull_request):
        # `PolicyViolation.created_at` is `auto_now_add`, so it can only be backdated with a
        # second `save(update_fields=...)`, never at creation time.
        violation = PolicyViolationFactory(pull_request=pull_request)
        violation.created_at = datetime.datetime(2026, 8, 12, tzinfo=datetime.UTC)
        violation.save(update_fields=["created_at"])
        return violation

    other_repository = RepositoryFactory()
    other_project = ProjectFactory()
    other_project.repositories.add(other_repository)
    other_identity = IdentityFactory(person=PersonFactory())
    other_pr = PullRequestFactory(
        repository=other_repository,
        author=other_identity,
        state="merged",
        created_at=datetime.datetime(2026, 8, 5, tzinfo=datetime.UTC),
        merged_at=datetime.datetime(2026, 8, 10, tzinfo=datetime.UTC),
        last_activity_at=datetime.datetime(2026, 8, 10, tzinfo=datetime.UTC),
    )
    _violation_in_period(other_pr)

    project = seeded_org["project"]
    repository = seeded_org["repository"]
    own_pr = seeded_org["people"][0][1]
    _violation_in_period(own_pr)
    rebuild(DATE_FROM, DATE_TO)
    bump_data_version()

    scope = Scope(scope_type=ScopeType.PROJECT, scope_id=project.id, access=ScopeFilter(unrestricted=True))
    user = django_user_model.objects.create_user(username="lead", password="pw")

    workbook = _workbook(scope, _params(), user=user)
    violations_sheet = workbook["Violations"]
    repo_names = {
        violations_sheet.cell(row=r, column=1).value for r in range(2, violations_sheet.max_row + 1)
    }
    assert repo_names == {repository.full_name}
    assert other_repository.full_name not in repo_names


def test_report_filename_is_ascii_and_scope_slugged():
    global_scope = Scope(scope_type=ScopeType.GLOBAL, scope_id=None, access=ScopeFilter(unrestricted=True))
    project_scope = Scope(scope_type=ScopeType.PROJECT, scope_id=7, access=ScopeFilter(unrestricted=True))

    period_name = report_filename(global_scope, _params())
    day_name = report_filename(project_scope, _params(mode="day", day=DATE_TO))

    assert period_name == "pr-radar_report_global_2026-08-01_2026-08-31.xlsx"
    assert day_name == "pr-radar_report_project-7_2026-08-31.xlsx"
    assert period_name.isascii()
    assert day_name.isascii()
