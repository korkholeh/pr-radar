"""T3/T4: a restricted lead (real `UserProjectAccess` grants, not the phase-1 stub) must see no
other project's data in a page, a chart JSON endpoint, a CSV export or an XLSX export — each case
paired with a positive assertion that their *own* project's data is present, so a broken page
cannot pass as "isolated" by returning nothing. T4 also checks the global level is exactly the sum
of the caller's accessible projects, and that an out-of-scope id is dropped silently in the query
string but 404s in the path."""

from __future__ import annotations

import datetime
import io

import openpyxl
import pytest
from django.urls import reverse

from apps.accounts.factories import UserProjectAccessFactory
from apps.activity.factories import PullRequestFactory
from apps.catalog.factories import IdentityFactory, PersonFactory, ProjectFactory, RepositoryFactory
from apps.metrics.rollups import rebuild
from apps.metrics.services import bump_data_version

DATE_FROM = datetime.date(2026, 8, 1)
DATE_TO = datetime.date(2026, 8, 31)
PERIOD_QS = "preset=custom&from=2026-08-01&to=2026-08-31&granularity=week"


def _seed_merged_prs(project, repository, count: int) -> None:
    for _ in range(count):
        identity = IdentityFactory(person=PersonFactory())
        PullRequestFactory(
            repository=repository,
            author=identity,
            state="merged",
            created_at=datetime.datetime(2026, 8, 5, tzinfo=datetime.UTC),
            merged_at=datetime.datetime(2026, 8, 10, tzinfo=datetime.UTC),
            last_activity_at=datetime.datetime(2026, 8, 10, tzinfo=datetime.UTC),
        )


def _project_with_repo(name: str):
    repository = RepositoryFactory(full_name=f"{name}/repo")
    project = ProjectFactory(name=name)
    project.repositories.add(repository)
    return project, repository


@pytest.fixture
def two_projects():
    own, own_repo = _project_with_repo("Own Project")
    other, other_repo = _project_with_repo("Other Project")
    _seed_merged_prs(own, own_repo, 2)
    _seed_merged_prs(other, other_repo, 3)
    rebuild(DATE_FROM, DATE_TO)
    bump_data_version()
    return own, other


@pytest.fixture
def restricted_lead(lead_user, two_projects):
    own, _other = two_projects
    UserProjectAccessFactory(user=lead_user, project=own)
    return lead_user


@pytest.mark.django_db
def test_restricted_lead_page_excludes_other_project(client, restricted_lead, two_projects):
    own, other = two_projects
    client.force_login(restricted_lead)

    response = client.get(reverse("dashboards:overview") + f"?{PERIOD_QS}")

    assert response.status_code == 200
    body = response.content.decode()
    assert own.name in body
    assert other.name not in body

    other_page = client.get(reverse("dashboards:project", args=[other.pk]) + f"?{PERIOD_QS}")
    assert other_page.status_code == 404

    own_page = client.get(reverse("dashboards:project", args=[own.pk]) + f"?{PERIOD_QS}")
    assert own_page.status_code == 200


@pytest.mark.django_db
def test_restricted_lead_chart_json_excludes_other_project(client, restricted_lead, two_projects):
    own, _other = two_projects
    client.force_login(restricted_lead)

    response = client.get(
        reverse("dashboards:chart_json", args=["throughput"]) + f"?{PERIOD_QS}",
        HTTP_ACCEPT="application/json",
    )

    assert response.status_code == 200
    body = response.json()
    totals = [sum(v for v in dataset["data"] if v is not None) for dataset in body["datasets"]]
    # Own project merged 2 PRs, the other project merged 3; a leaking global aggregate would sum
    # to (at least) 5. The restricted total must reflect only the caller's own project.
    assert sum(totals) == 2

    own_scoped = client.get(
        reverse("dashboards:chart_json", args=["throughput"])
        + f"?{PERIOD_QS}&scope_type=project&scope_id={own.pk}",
        HTTP_ACCEPT="application/json",
    )
    assert own_scoped.status_code == 200
    own_totals = [
        sum(v for v in dataset["data"] if v is not None) for dataset in own_scoped.json()["datasets"]
    ]
    assert sum(own_totals) == 2


@pytest.mark.django_db
def test_restricted_lead_csv_export_excludes_other_project(client, restricted_lead, two_projects):
    own, other = two_projects
    client.force_login(restricted_lead)

    response = client.get(reverse("dashboards:export", args=["projects", "csv"]) + f"?{PERIOD_QS}")

    assert response.status_code == 200
    body = b"".join(response.streaming_content).decode()
    assert own.name in body
    assert other.name not in body


@pytest.mark.django_db
def test_restricted_lead_xlsx_export_excludes_other_project(client, restricted_lead, two_projects):
    own, other = two_projects
    client.force_login(restricted_lead)

    response = client.get(reverse("dashboards:export", args=["projects", "xlsx"]) + f"?{PERIOD_QS}")

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    cell_values = {cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
    assert own.name in cell_values
    assert other.name not in cell_values


@pytest.mark.django_db
def test_restricted_lead_report_excludes_other_project(client, restricted_lead, two_projects):
    """T19's report-level case of the same criterion (T3 lists page/chart/CSV/XLSX; the report
    workbook is a fifth exit and gets its own case here since `build_report()` reads several
    selectors of its own, any one of which could leak)."""
    own, other = two_projects
    client.force_login(restricted_lead)

    response = client.get(reverse("dashboards:export_report") + f"?{PERIOD_QS}")

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    cell_values = {
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert own.name in cell_values
    assert other.name not in cell_values


@pytest.mark.django_db
def test_global_equals_sum_of_accessible_projects(client, lead_user):
    granted_a, repo_a = _project_with_repo("Granted A")
    granted_b, repo_b = _project_with_repo("Granted B")
    ungranted, repo_c = _project_with_repo("Ungranted")
    _seed_merged_prs(granted_a, repo_a, 2)
    _seed_merged_prs(granted_b, repo_b, 3)
    _seed_merged_prs(ungranted, repo_c, 10)
    rebuild(DATE_FROM, DATE_TO)
    bump_data_version()
    UserProjectAccessFactory(user=lead_user, project=granted_a)
    UserProjectAccessFactory(user=lead_user, project=granted_b)
    client.force_login(lead_user)

    response = client.get(
        reverse("dashboards:chart_json", args=["throughput"]) + f"?{PERIOD_QS}",
        HTTP_ACCEPT="application/json",
    )

    body = response.json()
    totals = [sum(v for v in dataset["data"] if v is not None) for dataset in body["datasets"]]
    assert sum(totals) == 5


@pytest.mark.django_db
def test_out_of_scope_query_ids_are_dropped_silently(client, restricted_lead, two_projects):
    own, other = two_projects
    client.force_login(restricted_lead)

    filtered = client.get(reverse("dashboards:overview") + f"?{PERIOD_QS}&project={other.pk}")

    # A validation error or a silent widening would both be wrong here: the out-of-scope id must
    # be dropped, leaving the page exactly as it is without the filter — the caller's own data
    # still renders (never emptied out) and the other project's name never appears.
    assert filtered.status_code == 200
    body = filtered.content.decode()
    assert own.name in body
    assert other.name not in body


@pytest.mark.django_db
def test_out_of_scope_path_id_is_404(client, restricted_lead, two_projects):
    _own, other = two_projects
    client.force_login(restricted_lead)

    response = client.get(reverse("dashboards:project", args=[other.pk]) + f"?{PERIOD_QS}")

    assert response.status_code == 404
